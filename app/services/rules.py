from dataclasses import dataclass
from pathlib import Path

from app.core.config import RULES_XLSX_PATH, SIX_ELEMENT_ORDER
from app.core.exceptions import RulesLoadError


@dataclass
class SixElementRule:
    name: str
    weight: str
    required_questions: str
    guidance: str
    target: str
    blockers: str
    strategy: str


def load_six_element_rules(xlsx_path: Path | None = None) -> dict[str, SixElementRule]:
    """Load six-element rules from the Excel source of truth."""
    target_path = xlsx_path or RULES_XLSX_PATH

    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as error:
        raise RulesLoadError("缺少 openpyxl。请先安装依赖后再运行。") from error

    if not target_path.exists():
        raise RulesLoadError(f"六要素规则文件不存在：{target_path}")

    workbook = load_workbook(target_path, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise RulesLoadError("六要素规则文件为空。")

    headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    header_map = {name: index for index, name in enumerate(headers)}
    required_headers = [
        "六要素",
        "权重",
        "必清事项（问题）",
        "参考信息指引",
        "目标",
        "典型卡点",
        "策略",
    ]
    missing_headers = [name for name in required_headers if name not in header_map]
    if missing_headers:
        raise RulesLoadError(f"六要素规则文件缺少表头：{', '.join(missing_headers)}")

    rules: dict[str, SixElementRule] = {}
    for row in rows[1:]:
        if not row or row[header_map["六要素"]] is None:
            continue

        name = str(row[header_map["六要素"]]).strip()
        if not name:
            continue

        rules[name] = SixElementRule(
            name=name,
            weight=str(row[header_map["权重"]] or "").strip(),
            required_questions=str(row[header_map["必清事项（问题）"]] or "").strip(),
            guidance=str(row[header_map["参考信息指引"]] or "").strip(),
            target=str(row[header_map["目标"]] or "").strip(),
            blockers=str(row[header_map["典型卡点"]] or "").strip(),
            strategy=str(row[header_map["策略"]] or "").strip(),
        )

    missing_elements = [name for name in SIX_ELEMENT_ORDER if name not in rules]
    if missing_elements:
        raise RulesLoadError(f"六要素规则缺失：{', '.join(missing_elements)}")

    return rules


def format_rules_context(rules: dict[str, SixElementRule]) -> str:
    """Render Excel rules into a prompt-friendly block."""
    blocks: list[str] = []
    for name in SIX_ELEMENT_ORDER:
        rule = rules[name]
        blocks.append(
            "\n".join(
                [
                    f"【{rule.name}】权重：{rule.weight}",
                    f"必清事项：{rule.required_questions}",
                    f"参考指引：{rule.guidance}",
                    f"目标：{rule.target}",
                    f"典型卡点：{rule.blockers}",
                    f"策略库：{rule.strategy}",
                ]
            )
        )
    return "\n\n".join(blocks)
