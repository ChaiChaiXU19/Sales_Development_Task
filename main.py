import getpass
import os
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

try:
    from dotenv import load_dotenv as _dotenv_loader
except ModuleNotFoundError:
    _dotenv_loader = None


RULES_XLSX_PATH = Path(__file__).resolve().parent / "data" / "six_elements_rules.xlsx"
SIX_ELEMENT_ORDER = [
    "需求",
    "技术认可",
    "决策链",
    "竞争对手",
    "合作伙伴",
    "流程",
]


@dataclass
class SixElementRule:
    name: str
    weight: str
    required_questions: str
    guidance: str
    target: str
    blockers: str
    strategy: str


def load_env_file_fallback(env_path: Path) -> None:
    """Fallback loader when python-dotenv is unavailable."""
    if not env_path.exists():
        return

    for raw_line in env_path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue

        key, value = line.split("=", 1)
        key = key.strip()
        if key.startswith("export "):
            key = key[len("export ") :].strip()
        value = value.strip().strip("'\"")
        if key:
            os.environ[key] = value


def write_api_key_to_local_env(env_path: Path, api_key: str) -> None:
    """Create or update .env with a local-only API key and restrictive permissions."""
    existing_lines: list[str] = []
    if env_path.exists():
        existing_lines = env_path.read_text(encoding="utf-8").splitlines()

    updated_lines: list[str] = []
    key_written = False

    for raw_line in existing_lines:
        stripped = raw_line.strip()
        normalized = stripped[len("export ") :].strip() if stripped.startswith("export ") else stripped

        if "=" in normalized:
            current_key = normalized.split("=", 1)[0].strip()
            if current_key == "OPENAI_API_KEY":
                updated_lines.append(f"OPENAI_API_KEY={api_key}")
                key_written = True
                continue

        updated_lines.append(raw_line)

    if not key_written:
        if updated_lines and updated_lines[-1].strip():
            updated_lines.append("")
        updated_lines.append(f"OPENAI_API_KEY={api_key}")

    env_path.write_text("\n".join(updated_lines).rstrip() + "\n", encoding="utf-8")
    env_path.chmod(0o600)


def prompt_and_persist_api_key(env_path: Path) -> str:
    """Prompt once for API key, save it to .env, and restrict file permissions."""
    try:
        api_key = getpass.getpass("请输入 OPENAI/DeepSeek API Key（输入内容不会显示）: ").strip()
    except (EOFError, KeyboardInterrupt) as error:
        raise ValueError("未获取到 API Key，已取消输入。") from error

    if not api_key:
        raise ValueError("API Key 不能为空。")

    try:
        write_api_key_to_local_env(env_path, api_key)
    except OSError as error:
        raise ValueError(f"无法写入本地 .env 文件：{error}") from error

    print("[INFO] API Key 已保存到本地 .env，且该文件不会进入 Git。")
    return api_key


def load_runtime_config() -> dict[str, Any]:
    """Load runtime configuration for DeepSeek and execution mode."""
    env_path = Path(__file__).resolve().parent / ".env"
    if _dotenv_loader is not None:
        _dotenv_loader(dotenv_path=env_path, override=True)
    else:
        load_env_file_fallback(env_path)

    api_key = (
        os.getenv("OPENAI_API_KEY", "").strip()
        or os.getenv("DEEPSEEK_API_KEY", "").strip()
        or os.getenv("API_KEY", "").strip()
        or os.getenv("api_key", "").strip()
    )
    api_base = (
        os.getenv("OPENAI_API_BASE", "").strip()
        or os.getenv("OPENAI_BASE_URL", "").strip()
        or os.getenv("DEEPSEEK_API_BASE", "").strip()
    )
    model_name = os.getenv("MODEL_NAME", "deepseek-chat").strip()
    debug_mode = os.getenv("BRAINSTORM_DEBUG", "false").lower() == "true"

    if not api_base:
        api_base = "https://api.deepseek.com"

    if not api_key and os.isatty(0):
        api_key = prompt_and_persist_api_key(env_path)

    if not api_key:
        raise ValueError(
            "缺少 OPENAI_API_KEY。请在项目根目录创建 .env 并配置，或在运行时输入 key。\n"
            "示例：OPENAI_API_KEY=你的key"
        )

    os.environ["OPENAI_API_KEY"] = api_key
    os.environ["OPENAI_API_BASE"] = api_base
    os.environ["OPENAI_BASE_URL"] = api_base
    os.environ.setdefault(
        "CREWAI_STORAGE_DIR",
        str((Path(__file__).resolve().parent / ".runtime" / "crewai").resolve()),
    )
    os.environ.setdefault("CREWAI_DISABLE_TELEMETRY", "true")
    os.environ.setdefault("CREWAI_DISABLE_TRACKING", "true")
    os.environ.setdefault("CREWAI_TRACING_ENABLED", "false")

    return {
        "api_key": api_key,
        "api_base": api_base,
        "model_name": model_name,
        "debug_mode": debug_mode,
    }


def load_six_element_rules(xlsx_path: Path) -> dict[str, SixElementRule]:
    """Load six-element rules from the Excel source of truth."""
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as error:
        raise ModuleNotFoundError(
            "缺少 openpyxl。请先安装依赖后再运行。"
        ) from error

    if not xlsx_path.exists():
        raise FileNotFoundError(f"六要素规则文件不存在：{xlsx_path}")

    workbook = load_workbook(xlsx_path, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("六要素规则文件为空。")

    headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    header_map = {name: index for index, name in enumerate(headers)}
    required_headers = [
        "六要素",
        "权重",
        "必清事项（问题）",
        "参考信息指引",
        "目标",
        "典型卡点",
        "策略",
    ]
    missing_headers = [name for name in required_headers if name not in header_map]
    if missing_headers:
        raise ValueError(f"六要素规则文件缺少表头：{', '.join(missing_headers)}")

    rules: dict[str, SixElementRule] = {}
    for row in rows[1:]:
        if not row or row[header_map["六要素"]] is None:
            continue

        name = str(row[header_map["六要素"]]).strip()
        if not name:
            continue

        rules[name] = SixElementRule(
            name=name,
            weight=str(row[header_map["权重"]] or "").strip(),
            required_questions=str(row[header_map["必清事项（问题）"]] or "").strip(),
            guidance=str(row[header_map["参考信息指引"]] or "").strip(),
            target=str(row[header_map["目标"]] or "").strip(),
            blockers=str(row[header_map["典型卡点"]] or "").strip(),
            strategy=str(row[header_map["策略"]] or "").strip(),
        )

    missing_elements = [name for name in SIX_ELEMENT_ORDER if name not in rules]
    if missing_elements:
        raise ValueError(f"六要素规则缺失：{', '.join(missing_elements)}")

    return rules


def format_rules_context(rules: dict[str, SixElementRule]) -> str:
    """Render Excel rules into a prompt-friendly block."""
    blocks: list[str] = []
    for name in SIX_ELEMENT_ORDER:
        rule = rules[name]
        blocks.append(
            "\n".join(
                [
                    f"【{rule.name}】权重：{rule.weight}",
                    f"必清事项：{rule.required_questions}",
                    f"参考指引：{rule.guidance}",
                    f"目标：{rule.target}",
                    f"典型卡点：{rule.blockers}",
                    f"策略库：{rule.strategy}",
                ]
            )
        )
    return "\n\n".join(blocks)


def format_project_context(
    project_name: str, six_element_inputs: dict[str, str]
) -> str:
    """Render project inputs into a stable prompt block."""
    blocks = [f"项目名称：{project_name}"]
    for name in SIX_ELEMENT_ORDER:
        blocks.append(f"【{name}】\n{six_element_inputs.get(name, '暂无输入（待补充）')}")
    return "\n\n".join(blocks)


def build_sales_battle_crew(shared_llm: Any, debug_mode: bool = False):
    """Build the six-element battle crew with a closing agent."""
    try:
        from crewai import Agent, Crew, Process, Task
    except ModuleNotFoundError as error:
        raise ModuleNotFoundError(
            "缺少 crewai。请先激活虚拟环境后安装依赖：\n"
            "source .venv/bin/activate && pip install crewai langchain_openai python-dotenv"
        ) from error

    diagnostician_agent = Agent(
        role="情报侦察兵（The Diagnostician）",
        goal="严格对照项目六要素标准，找出当前项目情报缺口、风险和高危卡点。",
        backstory=(
            "你是专门挑刺的销售情报侦察兵。你不负责想对策，只负责对照标准找盲区、"
            "识别高危卡点，并按风险高低排序。"
        ),
        llm=shared_llm,
        verbose=debug_mode,
        allow_delegation=False,
    )

    strategist_agent = Agent(
        role="破局军师（The Strategist）",
        goal="针对高危卡点，严格依据六要素策略库设计初版销售破局动作。",
        backstory=(
            "你是销售破局军师，擅长把项目弱点转化为下一步推进动作。"
            "你不能空谈，必须把动作写清楚：目标、执行动作、依赖资源、触达人和预期结果。"
        ),
        llm=shared_llm,
        verbose=debug_mode,
        allow_delegation=False,
    )

    challenger_agent = Agent(
        role="魔鬼代言人 / 蓝军（The Challenger）",
        goal="从客户领导或竞争对手视角推翻初版策略，指出时机、流程和竞争反制问题。",
        backstory=(
            "你是蓝军压力测试机。你的任务不是帮忙圆场，而是无情指出策略中的幼稚假设、"
            "流程错配、资源不足和竞争对手反击风险。"
        ),
        llm=shared_llm,
        verbose=debug_mode,
        allow_delegation=False,
    )

    closer_agent = Agent(
        role="销售主管 / The Closer",
        goal="从混乱讨论中提炼出销售真正要执行的 3-5 条动作。",
        backstory=(
            "你是只看结果的铁血销售主管。你砍掉一切华丽词藻，只保留可执行动作。"
            "你的唯一交付是 Who、When、Do What。"
        ),
        llm=shared_llm,
        verbose=debug_mode,
        allow_delegation=False,
    )

    task_1_diagnosis = Task(
        description=(
            "请基于以下项目现状和六要素规则，做一份高质量诊断报告。\n\n"
            "【项目现状】\n{project_context}\n\n"
            "【六要素规则】\n{rules_context}\n\n"
            "要求：\n"
            "1. 必须按六要素逐项输出。\n"
            "2. 每个要素固定包含：已知事实、关键缺口、风险等级、高危卡点。\n"
            "3. 只做诊断，不允许提出策略、建议、应对措施。\n"
            "4. 结合权重，明确当前最薄弱的要素和最危险的卡点。"
        ),
        expected_output=(
            "一份按六要素结构化输出的诊断报告，只包含事实、缺口、风险和卡点，不包含任何对策。"
        ),
        agent=diagnostician_agent,
    )

    task_2_strategy = Task(
        description=(
            "请只针对上一步识别出的高危卡点，输出初版销售破局动作。\n\n"
            "你必须同时参考以下规则库：\n{rules_context}\n\n"
            "要求：\n"
            "1. 只针对高危卡点出招，不要复述全部诊断。\n"
            "2. 每条动作必须包含：对应六要素、动作目标、执行动作、依赖资源、触达人、预期结果。\n"
            "3. 只能依据六要素策略库出招，不能凭空发挥。\n"
            "4. 动作必须可执行，避免抽象口号。"
        ),
        expected_output=(
            "一份初版破局动作清单，每条动作都包含对应要素、目标、动作、依赖资源、触达人和预期结果。"
        ),
        agent=strategist_agent,
        context=[task_1_diagnosis],
    )

    task_3_challenge = Task(
        description=(
            "请站在客户的降本增效领导或强势竞争对手视角，对上一步初版策略做压力测试。\n\n"
            "你必须输出：\n"
            "1. 哪些动作太天真或时机错误\n"
            "2. 哪些动作与当前采购流程冲突\n"
            "3. 竞争对手最可能的反制方式\n"
            "4. 必须修正的点\n\n"
            "要求：\n"
            "1. 必须逐条挑战初版动作，而不是泛泛评论。\n"
            "2. 不能生成最终方案，只能输出挑战意见和必须修正点。\n"
            "3. 挑战要结合六要素中的权重、流程阶段和竞争风险。"
        ),
        expected_output=(
            "一份蓝军挑战报告，包含对初版动作的逐条反驳、风险说明和必须修正点。"
        ),
        agent=challenger_agent,
        context=[task_1_diagnosis, task_2_strategy],
    )

    task_4_closer = Task(
        description=(
            "请旁听完前面所有历史讨论后，只输出最终收口动作清单。\n\n"
            "要求：\n"
            "1. 只提炼 3-5 条动作。\n"
            "2. 优先覆盖六要素中最薄弱、风险最高的环节。\n"
            "3. 每条动作必须固定格式：Who: ... | When: ... | Do What: ...\n"
            "4. Who 必须明确是销售本人、合作伙伴、技术人员、高层或其他角色。\n"
            "5. When 必须明确时间节点、最晚时间或阶段窗口。\n"
            "6. Do What 必须是可量化、可执行的动作，不允许空话。\n"
            "7. 禁止输出前言、结论、分析、解释，只输出最终动作清单。"
        ),
        expected_output=(
            "仅输出 3-5 条最终动作清单，每条使用 Who / When / Do What 格式，不包含其他说明。"
        ),
        agent=closer_agent,
        context=[task_1_diagnosis, task_2_strategy, task_3_challenge],
        markdown=True,
    )

    return Crew(
        agents=[
            diagnostician_agent,
            strategist_agent,
            challenger_agent,
            closer_agent,
        ],
        tasks=[
            task_1_diagnosis,
            task_2_strategy,
            task_3_challenge,
            task_4_closer,
        ],
        process=Process.sequential,
        verbose=debug_mode,
    )


def run_brainstorm_from_six_elements(
    project_name: str,
    six_element_inputs: dict[str, str],
    rules: dict[str, SixElementRule] | None = None,
) -> str | None:
    """Execute the sales battle plan from structured six-element inputs."""
    config = load_runtime_config()

    try:
        from langchain_openai import ChatOpenAI
    except ModuleNotFoundError as error:
        raise ModuleNotFoundError(
            "缺少 langchain_openai。请先激活虚拟环境后安装依赖：\n"
            "source .venv/bin/activate && pip install crewai langchain_openai python-dotenv"
        ) from error

    active_rules = rules or load_six_element_rules(RULES_XLSX_PATH)
    rules_context = format_rules_context(active_rules)
    project_context = format_project_context(project_name, six_element_inputs)

    print("[INFO] 已加载 .env 配置")
    print(f"[INFO] 使用模型: {config['model_name']}")
    print(f"[INFO] 使用 API Base: {config['api_base']}")
    print(f"[INFO] 已加载六要素规则文件: {RULES_XLSX_PATH}")

    shared_llm = ChatOpenAI(
        model=config["model_name"],
        api_key=config["api_key"],
        base_url=config["api_base"],
        temperature=0.4,
    )

    print("[INFO] 正在创建六要素销售作战 Crew（Diagnostician -> Strategist -> Challenger -> Closer）...")
    crew = build_sales_battle_crew(
        shared_llm=shared_llm,
        debug_mode=config["debug_mode"],
    )

    inputs = {
        "project_name": project_name,
        "project_context": project_context,
        "rules_context": rules_context,
    }

    print("[INFO] 开始生成销售下一步动作计划...")
    try:
        result = crew.kickoff(inputs=inputs)
    except Exception as error:
        print(f"[ERROR] 任务执行失败: {error}")
        return None

    if config["debug_mode"]:
        print("\n[DEBUG] 中间任务输出如下：")
        for index, task_output in enumerate(result.tasks_output, start=1):
            print(f"\n----- Task {index} Output Start -----")
            print(task_output.raw)
            print(f"----- Task {index} Output End -----")

    print("\n" + "=" * 80)
    print("最终《销售下一步动作计划》")
    print("=" * 80)
    print(result.raw)
    print("=" * 80)
    return result.raw


def ask_required_single_line(prompt_text: str) -> str:
    """Read a required single-line input."""
    while True:
        value = input(prompt_text).strip()
        if value:
            return value
        print("[WARN] 输入不能为空，请重新输入。")


def read_markdown_input_file(file_path: str) -> str:
    """Read a single markdown file that contains the whole project context."""
    path = Path(file_path).expanduser()
    if not path.is_absolute():
        path = (Path.cwd() / path).resolve()

    if not path.exists():
        raise ValueError(f"文件不存在：{path}")
    if not path.is_file():
        raise ValueError(f"不是普通文件：{path}")
    if path.suffix.lower() != ".md":
        raise ValueError("仅支持读取 .md 文件。")

    try:
        content = path.read_text(encoding="utf-8").strip()
    except UnicodeDecodeError as error:
        raise ValueError("Markdown 文件不是有效的 UTF-8 文本，请转换编码后重试。") from error

    if not content:
        raise ValueError("Markdown 文件内容为空，请重新提供。")

    return content


def parse_markdown_project_input(markdown_text: str, source_name: str) -> tuple[str, dict[str, str]]:
    """Parse one markdown document into project name and six-element inputs."""
    normalized = markdown_text.replace("\r\n", "\n").replace("\r", "\n")
    sections: dict[str, list[str]] = {name: [] for name in SIX_ELEMENT_ORDER}
    current_section: str | None = None
    project_name = ""

    project_name_patterns = [
        re.compile(r"^\s*项目名称\s*[:：]\s*(.+?)\s*$"),
        re.compile(r"^\s*#\s*项目名称\s*[:：]?\s*(.+?)\s*$"),
        re.compile(r"^\s*##\s*项目名称\s*[:：]?\s*(.+?)\s*$"),
    ]

    heading_pattern = re.compile(r"^\s{0,3}(#{1,6})\s*(.+?)\s*$")

    for line in normalized.split("\n"):
        stripped = line.strip()

        if not project_name:
            for pattern in project_name_patterns:
                match = pattern.match(stripped)
                if match:
                    project_name = match.group(1).strip()
                    break

        heading_match = heading_pattern.match(line)
        if heading_match:
            heading_text = heading_match.group(2).strip()
            matched_section = next(
                (name for name in SIX_ELEMENT_ORDER if heading_text == name),
                None,
            )
            if matched_section is not None:
                current_section = matched_section
                continue

        if current_section is not None:
            sections[current_section].append(line)

    if not project_name:
        first_heading = next(
            (
                match.group(2).strip()
                for raw_line in normalized.split("\n")
                if (match := heading_pattern.match(raw_line)) and match.group(1) == "#"
            ),
            "",
        )
        project_name = first_heading or Path(source_name).stem

    parsed_inputs = {
        name: "\n".join(lines).strip() or "暂无输入（待补充）"
        for name, lines in sections.items()
    }

    missing_sections = [name for name, content in parsed_inputs.items() if content == "暂无输入（待补充）"]
    if len(missing_sections) == len(SIX_ELEMENT_ORDER):
        raise ValueError(
            "Markdown 文件未识别到六要素内容。请使用二级标题，例如：## 需求、## 技术认可。"
        )

    return project_name, parsed_inputs


def load_project_input_from_markdown(file_path: str) -> tuple[str, dict[str, str]]:
    """Load and parse one markdown file as the only business input."""
    markdown_text = read_markdown_input_file(file_path)
    return parse_markdown_project_input(markdown_text, source_name=file_path)


if __name__ == "__main__":
    print("[INFO] 请提供一个包含项目名称和六要素内容的 Markdown 文件。")
    print("[INFO] 建议格式：# 项目名称 或 项目名称：xxx，然后使用 ## 需求 / ## 技术认可 / ... / ## 流程")
    try:
        markdown_file_path = ask_required_single_line("Markdown 文件路径：")
        project_name, six_element_inputs = load_project_input_from_markdown(markdown_file_path)
        print(f"[INFO] 已解析项目名称：{project_name}")
    except (KeyboardInterrupt, EOFError):
        print("\n[INFO] 用户取消输入，程序结束。")
        raise SystemExit(0)
    except ValueError as error:
        print(f"[ERROR] 输入文件解析失败: {error}")
        raise SystemExit(1)

    run_brainstorm_from_six_elements(
        project_name=project_name,
        six_element_inputs=six_element_inputs,
    )
